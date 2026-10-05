"""Per-file-type extraction into one page format.

Page record:
  {page_no, method: merge|surya|surya+tesseract|direct|blank, page_class, size: [w, h] (PDF points or None),
   blocks: [{id, label, kind: text|table|skipped, bbox, confidence, text, lines | table}],
   changes: [...], metrics: {...}, ocr: {mode, seconds}}

Labels follow Surya's vocabulary (SectionHeader, Text, Table, ListItem, PageHeader, PageFooter, ...)
so Word/Excel/text produce the same structure as PDFs.
"""
from __future__ import annotations

import csv
import io
import re
import shutil
import subprocess
from pathlib import Path

import pymupdf

from core.config import settings
from core.textnorm import clean
from ingest import fusion, surya_ocr, textlayer

PDF_TYPES = {".pdf"}
WORD_TYPES = {".docx"}
SHEET_TYPES = {".xlsx", ".xlsm", ".csv"}
TEXT_TYPES = {".txt", ".md", ".markdown"}
SUPPORTED = PDF_TYPES | WORD_TYPES | SHEET_TYPES | TEXT_TYPES


def file_kind(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in PDF_TYPES:
        return "pdf"
    if ext in WORD_TYPES:
        return "word"
    if ext in SHEET_TYPES:
        return "sheet"
    if ext in TEXT_TYPES:
        return "text"
    raise ValueError(f"unsupported file type: {ext} (supported: {', '.join(sorted(SUPPORTED))})")


# ------------------------------------------------------------------ PDF

def pdf_pages(path: Path, sha: str, pages: list[int] | None = None, progress=None):
    """Yield page records for a PDF. Pages are OCR'd in batches; results are cached per page."""
    doc = pymupdf.open(path)
    page_nos = pages or list(range(1, doc.page_count + 1))
    batch = settings.surya_batch
    for i in range(0, len(page_nos), batch):
        chunk = page_nos[i:i + batch]
        info = {}
        for p in chunk:
            page = doc[p - 1]
            ws = textlayer.words(page)
            info[p] = (page, ws, textlayer.classify(page, ws))
        to_ocr = {p: textlayer.render(info[p][0], settings.surya_dpi)
                  for p in chunk if info[p][2] != "blank" and not surya_ocr.cache_path(sha, p).exists()}
        cached = {p: None for p in chunk if info[p][2] != "blank"}
        ocr = surya_ocr.ocr_pages(sha, {**cached, **to_ocr}) if cached else {}
        for p in chunk:
            page, ws, cls = info[p]
            yield _pdf_page(p, page, ws, cls, ocr.get(p))
            if progress:
                progress(p, len(page_nos))


def _pdf_page(p: int, page, ws, cls: str, sp: dict | None) -> dict:
    size = [round(page.rect.width, 2), round(page.rect.height, 2)]
    base = {"page_no": p, "page_class": cls, "size": size}
    if cls == "blank" or sp is None:
        return {**base, "method": "blank", "blocks": [], "changes": [], "metrics": {"blocks": 0}, "ocr": None}
    if cls == "digital":
        method, ref, mode = "merge", ws, "textlayer"
    else:
        ref = tesseract_words(page)
        method, mode = ("surya+tesseract", "vote") if ref is not None else ("surya", "vote")
        ref = ref or []
    merged = fusion.merge_page(sp, ref, (page.rect.width, page.rect.height), mode=mode)
    confs = [b["confidence"] for b in sp["blocks"] if b.get("confidence") is not None]
    merged["metrics"].update(
        ocr_mode=sp["mode"], ocr_seconds=sp["seconds"],
        ocr_confidence=round(min(confs), 4) if confs else None,
        textlayer_words=len(ws), ocr_chars=sum(len(b.get("text", "")) for b in merged["blocks"]),
        textlayer_chars=sum(len(w.text) for w in ws),
    )
    if method == "surya":    # no second reader available: every number is unconfirmed
        for b in merged["blocks"]:
            for v in re.findall(r"\S*[0-9০-৯]\S*", b.get("text", "")):
                merged["changes"].append({"kind": "UNCONFIRMED", "value_kind": "num", "ocr": v, "ref": None,
                                          "block": b["id"], "loc": None, "bbox": b["bbox"],
                                          "reason": "no second OCR engine"})
        merged["metrics"]["unconfirmed"] = len(merged["changes"])
    return {**base, "method": method, **merged, "ocr": {"mode": sp["mode"], "seconds": sp["seconds"]}}


def tesseract_words(page, dpi: int = 300):
    """Word boxes from Tesseract (tessdata_best ben+eng) in PDF points, or None if not installed."""
    exe = shutil.which(settings.tesseract_cmd)
    if not exe or not (settings.tessdata_dir / "ben.traineddata").exists():
        return None
    img = page.get_pixmap(dpi=dpi).tobytes("png")
    out = subprocess.run([exe, "stdin", "stdout", "-l", "ben+eng", "--psm", "6",
                          "--tessdata-dir", str(settings.tessdata_dir), "tsv"],
                         input=img, capture_output=True, check=True).stdout.decode("utf-8")
    scale = 72 / dpi
    ws = []
    for row in csv.DictReader(io.StringIO(out), delimiter="\t", quoting=csv.QUOTE_NONE):
        text = clean((row.get("text") or "").strip())
        if not text or float(row.get("conf") or -1) < 0:
            continue
        x, y, w, h = (int(row[k]) * scale for k in ("left", "top", "width", "height"))
        ws.append(textlayer.Word(text, (x, y, x + w, y + h), "tesseract", False, int(row["line_num"])))
    return ws


# ------------------------------------------------------------------ Word / Excel / text

def _block(k, label, text=None, table=None):
    if table is not None:
        rows = [[clean(c) for c in r] for r in table]
        return {"id": k, "label": "Table", "kind": "table", "bbox": None, "confidence": None,
                "table": {"rows": rows, "header_rows": 1 if rows else 0},
                "text": "\n".join(" | ".join(r) for r in rows)}
    text = clean(text)
    return {"id": k, "label": label, "kind": "text", "bbox": None, "confidence": None,
            "lines": text.split("\n"), "text": text}


def _direct_page(p, blocks):
    return {"page_no": p, "page_class": "direct", "size": None, "method": "direct", "blocks": blocks,
            "changes": [], "metrics": {"blocks": len(blocks)}, "ocr": None}


def word_pages(path: Path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(str(path))
    blocks = []
    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(el, d)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            label = "SectionHeader" if style.startswith(("heading", "title")) else \
                "ListItem" if "list" in style else "Text"
            blocks.append(_block(len(blocks), label, text))
        elif tag == "tbl":
            t = Table(el, d)
            rows = [[c.text.strip() for c in r.cells] for r in t.rows]
            if rows:
                blocks.append(_block(len(blocks), None, table=rows))
    yield _direct_page(1, blocks)


def sheet_pages(path: Path):
    """One page per sheet; the first non-empty row is taken as the header."""
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        text = raw.decode("utf-8-sig", errors="replace")
        sheets = {path.stem: list(csv.reader(io.StringIO(text)))}
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = {ws.title: [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
                  for ws in wb.worksheets}
    for p, (name, rows) in enumerate(sheets.items(), start=1):
        rows = [r for r in rows if any(str(c).strip() for c in r)]
        if not rows:
            continue
        width = max(len(r) for r in rows)
        rows = [list(r) + [""] * (width - len(r)) for r in rows]
        yield _direct_page(p, [_block(0, "SectionHeader", name), _block(1, None, table=rows)])


def text_pages(path: Path):
    text = path.read_text("utf-8", errors="replace")
    blocks = []
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        label = "SectionHeader" if re.match(r"#{1,6}\s", para) else "Text"
        blocks.append(_block(len(blocks), label, re.sub(r"^#{1,6}\s+", "", para)))
    yield _direct_page(1, blocks)


def extract(path: Path, sha: str, pages=None, progress=None):
    kind = file_kind(path)
    if kind == "pdf":
        return pdf_pages(path, sha, pages, progress)
    if kind == "word":
        return word_pages(path)
    if kind == "sheet":
        return sheet_pages(path)
    return text_pages(path)
